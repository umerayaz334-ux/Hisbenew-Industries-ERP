"""Build Easyship Imperial shipment import workbooks from pending ERP orders."""

from __future__ import annotations

import io
import re
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

KG_TO_LB = 2.2046226218

TEMPLATE_PATH = (
    Path(__file__).resolve().parents[2]
    / "assets"
    / "easyship"
    / "Easyship_Shipments_Template_Imperial.xlsx"
)

TEMPLATE_SHEET = "Easyship_Template"

# Column indexes are 1-based to match the official Easyship Imperial template.
COL = {
    "choose_courier_by": 1,
    "shipping_insurance": 2,
    "taxes_duties": 3,
    "platform": 4,
    "platform_order_number": 5,
    "order_tags": 6,
    "receiver_name": 7,
    "receiver_phone": 8,
    "receiver_email": 9,
    "receiver_tax_id": 10,
    "address_line_1": 11,
    "address_line_2": 12,
    "postal_code": 13,
    "city": 14,
    "state": 15,
    "country": 16,
    "length_in": 17,
    "width_in": 18,
    "height_in": 19,
    "weight_lb": 20,
    "category": 21,
    "contains_liquid": 22,
    "contains_battery": 23,
    "description": 24,
    "country_of_origin": 25,
    "sku": 26,
    "platform_product_id": 27,
    "mpn": 28,
    "gtin": 29,
    "customs_value": 30,
    "customs_currency": 31,
    "quantity": 32,
    "buyer_notes": 45,
    "seller_notes": 46,
}

USA_ZIP_RE = re.compile(r"\b(\d{5})(?:-\d{4})?\b")
STATE_ZIP_RE = re.compile(
    r"^(?P<city>.+?)[,\s]+(?P<state>[A-Za-z]{2})\s+(?P<zip>\d{5})(?:-\d{4})?\s*$"
)
CITY_ZIP_RE = re.compile(r"^(?P<city>.+?)[,\s]+(?P<zip>\d{5})(?:-\d{4})?\s*$")

CATEGORY_ALIASES = {
    "accessory": "Accessory (no-battery)",
    "accessories": "Accessory (no-battery)",
    "bag": "Bags & Luggages",
    "bags": "Bags & Luggages",
    "luggage": "Bags & Luggages",
    "fashion": "Fashion",
    "apparel": "Fashion",
    "clothing": "Fashion",
    "leather": "Fashion",
    "home": "Home Decor",
    "decor": "Home Decor",
    "furniture": "Home Decor",
    "beauty": "Health & Beauty (non-liquid)",
    "jewelry": "Jewelry",
    "jewellery": "Jewelry",
    "pet": "Pet Accessory",
    "game": "Board Games",
    "toy": "Toys",
    "book": "Books & Collectibles",
}

DEFAULT_CATEGORY = "Accessory (no-battery)"
DEFAULT_COUNTRY_OF_ORIGIN = "Pakistan"
DEFAULT_COUNTRY = "United States"


def kg_to_lb(weight_kg: float | None) -> float:
    return round(float(weight_kg or 0) * KG_TO_LB, 3)


def map_easyship_category(raw: str | None) -> str:
    text = str(raw or "").strip()
    if not text:
        return DEFAULT_CATEGORY
    lowered = text.lower()
    for needle, mapped in CATEGORY_ALIASES.items():
        if needle in lowered:
            return mapped
    # Already an Easyship accepted value.
    if text in {
        "Accessory (no-battery)",
        "Accessory (with battery)",
        "Audio Video",
        "Bags & Luggages",
        "Board Games",
        "Books & Collectibles",
        "Cameras",
        "Computers & Laptops",
        "Crafts",
        "Documents",
        "Dry Food & Supplements",
        "Fashion",
        "Gaming",
        "Health & Beauty",
        "Health & Beauty (non-liquid)",
        "Home Appliances",
        "Home Decor",
        "Jewelry",
        "Mobile Phones",
        "Pet Accessory",
        "Sports & Leisure",
        "Toys",
    }:
        return text
    return DEFAULT_CATEGORY


def parse_shipping_address(address_text: str | None) -> dict:
    raw = str(address_text or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    if not raw:
        return {
            "line1": "",
            "line2": "",
            "city": "",
            "state": "",
            "postal_code": "",
            "country": DEFAULT_COUNTRY,
        }

    parts = [part.strip(" ,") for part in re.split(r"[\n;]+", raw) if part.strip(" ,")]
    country = DEFAULT_COUNTRY
    city = ""
    state = ""
    postal_code = ""

    # Peel country / city-state-zip from the bottom.
    while parts:
        tail = parts[-1]
        lowered = tail.lower()
        if lowered in {"usa", "u.s.a.", "u.s.", "us", "united states", "united states of america"}:
            country = "United States"
            parts.pop()
            continue
        if lowered in {"pakistan", "pk"}:
            country = "Pakistan"
            parts.pop()
            continue

        state_match = STATE_ZIP_RE.match(tail)
        if state_match:
            city = state_match.group("city").strip(" ,")
            state = state_match.group("state").upper()
            postal_code = state_match.group("zip")
            parts.pop()
            break

        city_match = CITY_ZIP_RE.match(tail)
        if city_match:
            city = city_match.group("city").strip(" ,")
            postal_code = city_match.group("zip")
            parts.pop()
            break

        zip_match = USA_ZIP_RE.search(tail)
        if zip_match and not postal_code:
            postal_code = zip_match.group(1)
            remainder = (tail[: zip_match.start()] + tail[zip_match.end() :]).strip(" ,")
            if remainder and not city:
                # Try "City ST" before the ZIP.
                city_state = re.match(r"^(?P<city>.+?)[,\s]+(?P<state>[A-Za-z]{2})$", remainder)
                if city_state:
                    city = city_state.group("city").strip(" ,")
                    state = city_state.group("state").upper()
                else:
                    city = remainder
            parts.pop()
            break
        break

    line1 = parts[0] if parts else raw.split("\n")[0].strip()
    line2 = parts[1] if len(parts) > 1 else ""
    if not postal_code:
        zip_match = USA_ZIP_RE.search(raw)
        if zip_match:
            postal_code = zip_match.group(1)

    return {
        "line1": line1[:180],
        "line2": line2[:180],
        "city": city[:80],
        "state": state[:40],
        "postal_code": postal_code,
        "country": country,
    }


def _copy_row_style(ws, source_row: int, target_row: int, max_col: int = 46) -> None:
    for col in range(1, max_col + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target.font = copy(source.font)
            target.border = copy(source.border)
            target.fill = copy(source.fill)
            target.number_format = source.number_format
            target.protection = copy(source.protection)
            target.alignment = copy(source.alignment)


def _receiver_name(order) -> str:
    customer = getattr(order, "customer", None)
    return (
        str(getattr(order, "import_shipping_name", None) or "").strip()
        or str(getattr(order, "import_contact_name", None) or "").strip()
        or (str(getattr(customer, "name", None) or "").strip() if customer else "")
        or "Customer"
    )


def _receiver_phone(order) -> str:
    customer = getattr(order, "customer", None)
    return (
        str(getattr(order, "import_contact_phone", None) or "").strip()
        or (str(getattr(customer, "phone", None) or "").strip() if customer else "")
        or ""
    )


def _receiver_email(order) -> str:
    customer = getattr(order, "customer", None)
    email = str(getattr(customer, "email", None) or "").strip() if customer else ""
    if email:
        return email
    order_no = str(getattr(order, "order_no", None) or getattr(order, "id", "order")).strip()
    safe = re.sub(r"[^A-Za-z0-9]+", "", order_no) or "order"
    return f"shipping+{safe}@hisbenew.local"


def _destination_address(order) -> str:
    customer = getattr(order, "customer", None)
    return (
        str(getattr(order, "import_shipping_address", None) or "").strip()
        or (str(getattr(customer, "shipping_address", None) or "").strip() if customer else "")
        or (str(getattr(customer, "address", None) or "").strip() if customer else "")
        or ""
    )


def _order_items(order) -> list:
    items = list(getattr(order, "items", None) or [])
    return [item for item in items if int(getattr(item, "quantity", 0) or 0) > 0]


def _item_weight_kg(item) -> float:
    product = getattr(item, "product", None)
    unit = float(getattr(product, "unit_weight_kg", 0) or 0) if product else 0.0
    qty = max(int(getattr(item, "quantity", 0) or 0), 0)
    return round(unit * qty, 4)


def _item_customs_value(item, order) -> float:
    line_total = float(getattr(item, "line_total", 0) or 0)
    if line_total > 0:
        return round(line_total, 2)
    unit_price = float(getattr(item, "unit_price", 0) or 0)
    qty = max(int(getattr(item, "quantity", 0) or 0), 1)
    if unit_price > 0:
        return round(unit_price * qty, 2)
    order_total = float(getattr(order, "order_total_usd", 0) or getattr(order, "total_amount", 0) or 0)
    return round(max(order_total, 1.0), 2)


def _distribute_package_weight_lb(items: list, package_weight_lb: float) -> list[float]:
    weights_kg = [_item_weight_kg(item) for item in items]
    total_kg = sum(weights_kg)
    package_lb = max(float(package_weight_lb or 0), 0.001)
    if total_kg > 0:
        return [round(package_lb * (w / total_kg), 3) for w in weights_kg]
    share = round(package_lb / max(len(items), 1), 3)
    values = [share] * len(items)
    # Keep the sum close to the package weight on the first line.
    if values:
        values[0] = round(package_lb - sum(values[1:]), 3)
    return values


def build_easyship_workbook_bytes(
    orders: list,
    packages_by_order_id: dict[int, dict],
    *,
    choose_courier_by: str = "Best Value For Money",
    taxes_duties_paid_by: str = "Receiver",
    shipping_insurance: str = "No",
) -> bytes:
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"Easyship template missing: {TEMPLATE_PATH}")

    workbook = load_workbook(TEMPLATE_PATH)
    ws = workbook[TEMPLATE_SHEET]

    # Remove sample data rows while keeping header + validations.
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    row_index = 2
    for order in orders:
        order_id = int(getattr(order, "id"))
        package = packages_by_order_id.get(order_id) or {}
        length_in = float(package.get("length_in") or 0)
        width_in = float(package.get("width_in") or 0)
        height_in = float(package.get("height_in") or 0)
        weight_lb = float(package.get("weight_lb") or 0)
        if length_in <= 0 or width_in <= 0 or height_in <= 0 or weight_lb <= 0:
            raise ValueError(
                f"Order #{getattr(order, 'order_no', order_id)} needs package length, width, height, and weight."
            )

        items = _order_items(order)
        if not items:
            # Still export a placeholder row so the shipment can be created.
            items = [None]

        address = parse_shipping_address(_destination_address(order))
        item_weights = _distribute_package_weight_lb(
            [item for item in items if item is not None] or [object()],
            weight_lb,
        )
        if items == [None]:
            item_weights = [weight_lb]

        for item_offset, item in enumerate(items):
            if row_index > 2:
                _copy_row_style(ws, 2, row_index)

            product = getattr(item, "product", None) if item is not None else None
            article_no = (
                str(getattr(product, "article_no", None) or "").strip()
                if product
                else ""
            )
            product_name = (
                str(getattr(product, "name", None) or "").strip()
                if product
                else "Order items"
            )
            category = map_easyship_category(
                getattr(product, "category", None) if product else None
            )
            quantity = max(int(getattr(item, "quantity", 1) or 1), 1) if item is not None else 1
            customs_value = (
                _item_customs_value(item, order)
                if item is not None
                else round(float(getattr(order, "order_total_usd", 0) or getattr(order, "total_amount", 1) or 1), 2)
            )

            values = {
                COL["choose_courier_by"]: choose_courier_by if item_offset == 0 else None,
                COL["shipping_insurance"]: shipping_insurance,
                COL["taxes_duties"]: taxes_duties_paid_by,
                COL["platform"]: str(getattr(order, "platform", None) or "Direct Sales").strip() or "Direct Sales",
                COL["platform_order_number"]: str(getattr(order, "order_no", "") or ""),
                COL["order_tags"]: "",
                COL["receiver_name"]: _receiver_name(order),
                COL["receiver_phone"]: _receiver_phone(order),
                COL["receiver_email"]: _receiver_email(order),
                COL["receiver_tax_id"]: "",
                COL["address_line_1"]: address["line1"] or "Address required",
                COL["address_line_2"]: address["line2"],
                COL["postal_code"]: address["postal_code"] or "",
                COL["city"]: address["city"] or "",
                COL["state"]: address["state"] or "",
                COL["country"]: address["country"] or DEFAULT_COUNTRY,
                COL["length_in"]: length_in,
                COL["width_in"]: width_in,
                COL["height_in"]: height_in,
                COL["weight_lb"]: item_weights[item_offset] if item_offset < len(item_weights) else weight_lb,
                COL["category"]: category,
                COL["contains_liquid"]: "No",
                COL["contains_battery"]: "No",
                COL["description"]: product_name or article_no or f"Order {getattr(order, 'order_no', '')}",
                COL["country_of_origin"]: DEFAULT_COUNTRY_OF_ORIGIN,
                COL["sku"]: article_no,
                COL["platform_product_id"]: "",
                COL["mpn"]: "",
                COL["gtin"]: "",
                COL["customs_value"]: max(customs_value, 0.01),
                COL["customs_currency"]: "USD",
                COL["quantity"]: quantity,
                COL["buyer_notes"]: "",
                COL["seller_notes"]: f"ERP order #{getattr(order, 'order_no', order_id)}",
            }
            for col_index, value in values.items():
                ws.cell(row_index, col_index, value)
            row_index += 1

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
